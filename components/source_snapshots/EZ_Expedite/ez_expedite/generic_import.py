from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .db import ensure_checklist_items, log_activity, next_case_number, set_custom_value, utcnow
from .importers import clean, iso_date, normalize_status

CORE_FIELDS = [
    "title",
    "description",
    "source",
    "customer",
    "supplier",
    "part_number",
    "revision",
    "work_order",
    "sales_order",
    "purchase_order",
    "department",
    "work_center",
    "owner_name",
    "owner_email",
    "priority",
    "status",
    "next_action",
    "due_date",
    "created_date",
]

ALIASES = {
    "title": ["title", "subject", "issue", "occurrence", "event", "summary"],
    "description": ["description", "details", "problem", "discrepancy", "notes", "comment"],
    "source": ["source", "origin"],
    "customer": ["customer", "customer name", "cust"],
    "supplier": ["supplier", "vendor", "vendor name"],
    "part_number": ["part number", "part no", "part #", "item", "item number"],
    "revision": ["revision", "rev"],
    "work_order": ["work order", "wo", "w.o.", "job", "job number"],
    "sales_order": ["sales order", "so", "s.o.", "order number"],
    "purchase_order": ["purchase order", "po", "p.o.", "customer po"],
    "department": ["department", "dept"],
    "work_center": ["work center", "workcentre", "wc"],
    "owner_name": ["owner", "responsible", "responsible employee", "assigned to"],
    "owner_email": ["owner email", "email", "assigned email"],
    "priority": ["priority", "severity"],
    "status": ["status", "state"],
    "next_action": ["next action", "action required", "required action"],
    "due_date": ["due date", "required date", "target date"],
    "created_date": ["create date", "created date", "date", "opened date"],
}


def _rows_xlsx(path: Path) -> tuple[list[str], list[list[Any]]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    iterator = ws.iter_rows(values_only=True)
    first = next(iterator, None)
    if not first:
        return [], []
    headers = [clean(v) for v in first]
    data = [list(row) for row in iterator if any(v not in (None, "") for v in row)]
    wb.close()
    return headers, data


def _rows_csv(path: Path) -> tuple[list[str], list[list[Any]]]:
    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        reader = csv.reader(handle)
        headers = [clean(v) for v in next(reader, [])]
        return headers, [row for row in reader if any(clean(v) for v in row)]


def read_tabular(path: str | Path) -> tuple[list[str], list[list[Any]]]:
    path = Path(path)
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return _rows_xlsx(path)
    if path.suffix.lower() == ".csv":
        return _rows_csv(path)
    raise ValueError("Supported generic imports are .xlsx, .xlsm, and .csv")


def auto_mapping(headers: list[str]) -> dict[str, str]:
    normalized = {h.lower().strip(): h for h in headers if h}
    mapping: dict[str, str] = {}
    for target, aliases in ALIASES.items():
        for alias in aliases:
            if alias in normalized:
                mapping[target] = normalized[alias]
                break
    return mapping


def import_generic_file(
    con,
    path: str | Path,
    type_id: int,
    mapping: dict[str, str],
    external_system: str = "",
    external_id_column: str = "",
    actor: str = "Generic import",
) -> dict[str, int]:
    headers, rows = read_tabular(path)
    if not headers:
        raise ValueError("The import file has no header row.")
    header_index = {header: i for i, header in enumerate(headers)}
    trow = con.execute("SELECT name FROM occurrence_types WHERE id=? AND active=1", (type_id,)).fetchone()
    if not trow:
        raise ValueError("Occurrence type not found or inactive.")

    custom_defs = {
        int(row["id"]): row
        for row in con.execute(
            "SELECT * FROM custom_field_defs WHERE occurrence_type_id=? AND active=1",
            (type_id,),
        )
    }
    created = updated = skipped = 0

    def value(source_header: str, row: list[Any]) -> Any:
        idx = header_index.get(source_header)
        return row[idx] if idx is not None and idx < len(row) else None

    for source_row in rows:
        external_id = clean(value(external_id_column, source_row)) if external_id_column else ""
        occurrence_id = None
        if external_system and external_id:
            found = con.execute(
                """SELECT occurrence_id FROM external_refs
                   WHERE system_name=? AND entity_type='IMPORT' AND external_id=?
                   ORDER BY id LIMIT 1""",
                (external_system, external_id),
            ).fetchone()
            if found:
                occurrence_id = int(found["occurrence_id"])

        data: dict[str, Any] = {}
        for target in CORE_FIELDS:
            source_header = mapping.get(target, "")
            if source_header:
                data[target] = value(source_header, source_row)

        title = clean(data.get("title"))
        description = clean(data.get("description"))
        if not title:
            title = description[:120] if description else ""
        if not title:
            title = f"{trow['name']} imported occurrence"

        now = utcnow()
        due = iso_date(data.get("due_date"))
        created_date = iso_date(data.get("created_date"))
        status = normalize_status(data.get("status"))
        priority = clean(data.get("priority")) or "Normal"

        values = (
            title,
            description,
            clean(data.get("source")) or external_system or "Spreadsheet import",
            clean(data.get("customer")),
            clean(data.get("supplier")),
            clean(data.get("part_number")),
            clean(data.get("revision")),
            clean(data.get("work_order")),
            clean(data.get("sales_order")),
            clean(data.get("purchase_order")),
            clean(data.get("department")),
            clean(data.get("work_center")),
            clean(data.get("owner_name")),
            clean(data.get("owner_email")),
            priority,
            status,
            clean(data.get("next_action")),
            due,
            created_date,
            now,
        )

        if occurrence_id:
            con.execute(
                """UPDATE occurrences SET
                   title=?,description=?,source=?,customer=?,supplier=?,part_number=?,revision=?,
                   work_order=?,sales_order=?,purchase_order=?,department=?,work_center=?,owner_name=?,owner_email=?,
                   priority=?,status=?,next_action=?,due_date=?,created_date=COALESCE(?,created_date),updated_at=?
                   WHERE id=?""",
                values + (occurrence_id,),
            )
            updated += 1
        else:
            case_number = next_case_number(con, type_id)
            cur = con.execute(
                """INSERT INTO occurrences(
                   case_number,occurrence_type_id,title,description,source,customer,supplier,part_number,revision,
                   work_order,sales_order,purchase_order,department,work_center,owner_name,owner_email,priority,status,
                   next_action,due_date,created_date,last_activity_at,created_at,updated_at)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (case_number, type_id) + values[:-1] + (now, now, now),
            )
            occurrence_id = int(cur.lastrowid)
            created += 1
            if trow["name"] == "RMA":
                con.execute(
                    "INSERT OR IGNORE INTO rma_details(occurrence_id,recovery_status) VALUES(?,?)",
                    (occurrence_id, "NOT REQUIRED"),
                )
            if external_system and external_id:
                con.execute(
                    """INSERT OR IGNORE INTO external_refs(
                       occurrence_id,system_name,entity_type,external_id,created_at)
                       VALUES(?,?,?,?,?)""",
                    (occurrence_id, external_system, "IMPORT", external_id, now),
                )

        ensure_checklist_items(con, occurrence_id, type_id)
        for target, source_header in mapping.items():
            if not target.startswith("custom:") or not source_header:
                continue
            try:
                field_id = int(target.split(":", 1)[1])
            except ValueError:
                continue
            if field_id not in custom_defs:
                continue
            set_custom_value(con, occurrence_id, field_id, clean(value(source_header, source_row)))

        log_activity(
            con,
            occurrence_id,
            "IMPORT",
            f"{trow['name']} synchronized from {external_system or Path(path).name}.",
            actor,
        )

    return {"created": created, "updated": updated, "skipped": skipped}
