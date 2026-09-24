from __future__ import annotations

import csv
import hashlib
import re
import shutil
from datetime import date, datetime
from pathlib import Path

from flask import jsonify, redirect, request
from openpyxl import load_workbook
from werkzeug.utils import secure_filename

from audit_journal import record_event
from forgeqc_app import Customer, WorkOrder, as_float, as_int, db, get_or_create, page, parse_date
from runtime_paths import erp_archive_dir, erp_inbox_dir


class ERPImportBatch(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    imported_at = db.Column(db.DateTime, default=datetime.utcnow, index=True)
    source_filename = db.Column(db.String(300), nullable=False)
    source_sha256 = db.Column(db.String(64), nullable=False, unique=True, index=True)
    report_type = db.Column(db.String(80), default="Unknown")
    source_row_count = db.Column(db.Integer, default=0)
    imported_row_count = db.Column(db.Integer, default=0)
    status = db.Column(db.String(80), default="Pending")
    notes = db.Column(db.Text)
    archived_path = db.Column(db.String(700))


class ERPShipment(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    import_batch_id = db.Column(db.Integer, db.ForeignKey("erp_import_batch.id"), nullable=False, index=True)
    ship_date = db.Column(db.Date, index=True)
    customer_name = db.Column(db.String(220))
    part_number = db.Column(db.String(180), index=True)
    work_order_number = db.Column(db.String(140), index=True)
    quantity_shipped = db.Column(db.Float, default=0)

    import_batch = db.relationship("ERPImportBatch")


ALIASES = {
    "work_order_number": {"work order", "work order number", "wo", "job", "job number", "workorder"},
    "customer": {"customer", "customer name", "cust name"},
    "part_number": {"part", "part number", "item", "item number", "part no"},
    "description": {"description", "part description", "item description"},
    "quantity_ordered": {"qty ordered", "quantity ordered", "order qty", "ordered quantity"},
    "quantity_completed": {"qty completed", "quantity completed", "completed qty", "complete qty"},
    "quantity_failed": {"qty failed", "quantity failed", "failed qty", "reject qty", "rejected qty"},
    "release_date": {"release date", "released", "start date"},
    "due_date": {"due date", "required date", "need date"},
    "status": {"status", "job status", "wo status"},
    "priority": {"priority"},
    "material_status": {"material status"},
    "ship_date": {"ship date", "shipped date", "invoice date", "date shipped"},
    "quantity_shipped": {"qty shipped", "quantity shipped", "shipped qty", "ship qty"},
}


def _norm(value):
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").strip().lower()).strip()


def _canonical_key(header):
    normalized = _norm(header)
    for key, aliases in ALIASES.items():
        if normalized == _norm(key) or normalized in aliases:
            return key
    return normalized.replace(" ", "_")


def _sha256(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _rows_from_csv(path):
    with Path(path).open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def _rows_from_excel(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    try:
        headers = [str(value or "").strip() for value in next(iterator)]
    except StopIteration:
        return []
    rows = []
    for values in iterator:
        rows.append({headers[i]: values[i] if i < len(values) else None for i in range(len(headers))})
    return rows


def read_rows(path):
    suffix = Path(path).suffix.lower()
    if suffix == ".csv":
        rows = _rows_from_csv(path)
    elif suffix in {".xlsx", ".xlsm"}:
        rows = _rows_from_excel(path)
    else:
        raise ValueError(f"Unsupported ERP report type: {suffix}")
    normalized = []
    for row in rows:
        normalized.append({_canonical_key(k): v for k, v in row.items()})
    return normalized


def detect_report_type(rows):
    keys = set()
    for row in rows[:20]:
        keys.update(row.keys())
    if "quantity_shipped" in keys and ("ship_date" in keys or "part_number" in keys):
        return "Shipment"
    if "work_order_number" in keys and "part_number" in keys:
        return "Work Order"
    return "Staged / Unmapped"


def _archive(path, digest):
    source = Path(path)
    destination = erp_archive_dir() / f"{datetime.utcnow().strftime('%Y%m%dT%H%M%SZ')}_{digest[:10]}_{secure_filename(source.name)}"
    shutil.move(str(source), str(destination))
    return destination


def ingest_file(path):
    source = Path(path)
    digest = _sha256(source)
    existing = ERPImportBatch.query.filter_by(source_sha256=digest).first()
    if existing:
        try:
            source.unlink()
        except OSError:
            pass
        record_event("ERP_IMPORT", "Duplicate skipped", entity_type="ERP_IMPORT", entity_id=existing.id, detail=source.name, data={"sha256": digest})
        return existing, False

    rows = read_rows(source)
    report_type = detect_report_type(rows)
    batch = ERPImportBatch(
        source_filename=source.name,
        source_sha256=digest,
        report_type=report_type,
        source_row_count=len(rows),
        status="Processing",
    )
    db.session.add(batch)
    db.session.flush()

    imported = 0
    notes = []
    try:
        if report_type == "Shipment":
            for row in rows:
                qty = as_float(row.get("quantity_shipped"), 0)
                if qty <= 0:
                    continue
                shipment = ERPShipment(
                    import_batch_id=batch.id,
                    ship_date=parse_date(row.get("ship_date")) or date.today(),
                    customer_name=str(row.get("customer") or "").strip(),
                    part_number=str(row.get("part_number") or "").strip(),
                    work_order_number=str(row.get("work_order_number") or "").strip(),
                    quantity_shipped=qty,
                )
                db.session.add(shipment)
                imported += 1
        elif report_type == "Work Order":
            for row in rows:
                wo_number = str(row.get("work_order_number") or "").strip()
                part_number = str(row.get("part_number") or "").strip()
                if not wo_number or not part_number:
                    continue
                work_order = WorkOrder.query.filter_by(work_order_number=wo_number).first() or WorkOrder(work_order_number=wo_number)
                customer = get_or_create(Customer, row.get("customer"))
                work_order.customer_id = customer.id if customer else work_order.customer_id
                work_order.part_number = part_number
                work_order.description = str(row.get("description") or work_order.description or "")
                work_order.quantity_ordered = as_int(row.get("quantity_ordered"), work_order.quantity_ordered or 0)
                work_order.quantity_completed = as_int(row.get("quantity_completed"), work_order.quantity_completed or 0)
                work_order.quantity_failed = as_int(row.get("quantity_failed"), work_order.quantity_failed or 0)
                work_order.release_date = parse_date(row.get("release_date")) or work_order.release_date
                work_order.due_date = parse_date(row.get("due_date")) or work_order.due_date
                work_order.status = str(row.get("status") or work_order.status or "Open")
                work_order.priority = str(row.get("priority") or work_order.priority or "Normal")
                work_order.material_status = str(row.get("material_status") or work_order.material_status or "Not Reviewed")
                db.session.add(work_order)
                imported += 1
        else:
            notes.append("Report archived for mapping review. No recognized Work Order or Shipment schema was detected.")

        destination = _archive(source, digest)
        batch.archived_path = str(destination)
        batch.imported_row_count = imported
        batch.status = "Imported" if report_type != "Staged / Unmapped" else "Staged / Unmapped"
        batch.notes = " ".join(notes)
        db.session.flush()
        record_event(
            "ERP_IMPORT",
            batch.status,
            entity_type="ERP_IMPORT",
            entity_id=batch.id,
            detail=batch.source_filename,
            data={
                "sha256": digest,
                "report_type": report_type,
                "source_rows": len(rows),
                "imported_rows": imported,
                "archive": str(destination),
            },
        )
        return batch, True
    except Exception as exc:
        batch.status = "Failed"
        batch.notes = str(exc)[:2000]
        record_event("ERP_IMPORT", "Failed", entity_type="ERP_IMPORT", entity_id=batch.id, detail=batch.source_filename, data={"error": str(exc)[:1000]})
        raise


def process_inbox():
    processed = []
    allowed = {".csv", ".xlsx", ".xlsm"}
    for path in sorted(erp_inbox_dir().iterdir()):
        if not path.is_file() or path.suffix.lower() not in allowed:
            continue
        batch, created = ingest_file(path)
        processed.append((batch.id, created, batch.status))
    if processed:
        db.session.commit()
        try:
            from pulse_intelligence import compute_pulse_snapshot
            compute_pulse_snapshot(30)
        except Exception:
            pass
    return processed


def register(app):
    @app.route("/erp-import", methods=["GET", "POST"])
    def erp_import_page():
        if request.method == "POST":
            uploaded = request.files.get("report")
            if uploaded and uploaded.filename:
                name = secure_filename(uploaded.filename)
                destination = erp_inbox_dir() / name
                uploaded.save(destination)
                record_event("ERP_DROP", "Uploaded", entity_type="ERP_FILE", entity_id=name, detail="ERP report placed in server inbox")
                process_inbox()
            return redirect("/erp-import")

        rows = "".join(
            f"<tr><td>{batch.imported_at}</td><td>{batch.source_filename}</td><td>{batch.report_type}</td><td>{batch.source_row_count}</td><td>{batch.imported_row_count}</td><td>{batch.status}</td><td>{batch.notes or ''}</td></tr>"
            for batch in ERPImportBatch.query.order_by(ERPImportBatch.id.desc()).limit(100)
        )
        body = f"""
        <div class='notice'>Drop custom ERP CSV/XLSX reports into <b>{erp_inbox_dir()}</b> or upload them here. The server hashes, archives, imports recognized data, and journals every ingest. Shipment reports become the preferred customer-PPM denominator.</div>
        <section><form method='post' enctype='multipart/form-data' class='form'>
          <label class='wide'>ERP Report<input type='file' name='report' accept='.csv,.xlsx,.xlsm' required></label>
          <div class='wide'><button>Upload + Process</button></div>
        </form></section><br>
        <section><table><tr><th>Imported</th><th>File</th><th>Type</th><th>Rows</th><th>Applied</th><th>Status</th><th>Notes</th></tr>{rows}</table></section>
        """
        return page("ERP Report Imports", body)

    @app.route("/erp-import/run", methods=["POST"])
    def erp_import_run():
        process_inbox()
        return redirect("/erp-import")

    @app.route("/api/erp/status")
    def erp_status():
        latest = ERPImportBatch.query.order_by(ERPImportBatch.id.desc()).first()
        return jsonify({
            "inbox": str(erp_inbox_dir()),
            "latest": {
                "id": latest.id,
                "filename": latest.source_filename,
                "report_type": latest.report_type,
                "status": latest.status,
                "imported_at": latest.imported_at.isoformat() if latest and latest.imported_at else None,
            } if latest else None,
        })
