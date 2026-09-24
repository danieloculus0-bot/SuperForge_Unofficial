from __future__ import annotations

import copy
import re
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

EXACT_R3_SHEET = "FAI FORM"
R3_START_ROW = 24
R3_TEMPLATE_END_ROW = 48
R3_TEMPLATE_CAPACITY = R3_TEMPLATE_END_ROW - R3_START_ROW + 1
R3_BODY_ROW_HEIGHT = 17.0
R3_HEADER_ROW_HEIGHT = 15.0

R3_COLUMNS = {
    "Char Number": 1,
    "Reference Location": 2,
    "Requirement LSL": 3,
    "Requirement Nominal": 4,
    "Requirement USL": 5,
    "Type": 6,
    "Supplier Actual": 7,
    "Supplier Yes": 8,
    "Supplier No": 9,
    "EZ Fabricating Actual": 10,
    "In Spec": 11,
    "Tooling Used": 13,
    "Comments": 14,
}

GENERIC_HEADERS = [
    "Char Number", "Reference Location", "Requirement LSL", "Requirement Nominal", "Requirement USL",
    "Type", "EZ Fabricating Actual", "In Spec", "Tooling Used", "Comments",
]

HEADER_ALIASES = {
    "Char Number": ["char", "char number", "char no", "characteristic", "balloon", "balloon no"],
    "Reference Location": ["reference location", "ref location", "location", "zone"],
    "Requirement LSL": ["requirement lsl", "lsl", "lower spec limit", "lower limit", "minimum", "min"],
    "Requirement Nominal": ["requirement nominal", "nominal", "requirement", "dimension", "specified requirement"],
    "Requirement USL": ["requirement usl", "usl", "upper spec limit", "upper limit", "maximum", "max"],
    "Type": ["type", "characteristic type", "dim type"],
    "EZ Fabricating Actual": ["ez fabricating actual", "actual", "actual result", "measured actual", "measurement"],
    "In Spec": ["in spec", "inspec", "pass fail", "pass/fail", "accept", "result"],
    "Tooling Used": ["tooling used", "tooling", "tool used", "gage", "gauge", "inspection tool"],
    "Comments": ["comments", "comment", "notes", "remark", "remarks"],
}

LIST_SHEET_SPECS = {
    "CHARACTERISTICS": [
        "NOTE", "TOLERANCE", "FINISH", "WELD", "MATERIAL", "RADIUS", "LINEAR", "DIAMETER", "Ø", "▱", "◯", "⌭",
        "∩", "⌓", "∠", "⊥", "//", "⌖", "◎", "⌯", "⌰", "Ⓕ", "Ⓛ", "°", "±",
        "↧", "≥", "≤", "⌴", "⌵", "µ", "✓", "℄",
    ],
    "TOOLING": ["VISUAL", "CALIPER", "CERTIFICATION", "TAPE ", "PROTRACTOR", "ANGLE GAGE", "THREAD GAGE", "HARDWARE", "FITMENT/NHA", "MICROMETER"],
    "ATTRIBUTE": ["PASS", "FAIL"],
}


def _norm(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def _safe_cell_value(characteristic: Any, attr: str, default: Any = "") -> Any:
    return getattr(characteristic, attr, default)


def _characteristic_metadata(characteristic: Any) -> dict[str, Any]:
    metadata = getattr(characteristic, "metadata", {}) or {}
    return metadata if isinstance(metadata, dict) else {}


def _row_values(characteristic: Any) -> dict[str, Any]:
    return {
        "Char Number": _safe_cell_value(characteristic, "char_number"),
        "Reference Location": _safe_cell_value(characteristic, "reference_location"),
        "Requirement LSL": _safe_cell_value(characteristic, "lsl"),
        "Requirement Nominal": _safe_cell_value(characteristic, "nominal"),
        "Requirement USL": _safe_cell_value(characteristic, "usl"),
        "Type": _safe_cell_value(characteristic, "type"),
        "EZ Fabricating Actual": _safe_cell_value(characteristic, "actual", ""),
        "Tooling Used": _safe_cell_value(characteristic, "tooling", ""),
        "Comments": _safe_cell_value(characteristic, "comments", ""),
    }


def _pick_sheet(wb):
    return wb[EXACT_R3_SHEET] if EXACT_R3_SHEET in wb.sheetnames else wb.active


def _is_r3_form(ws) -> bool:
    markers = [str(ws[cell].value or "").upper() for cell in ("A3", "A22", "C22", "G22", "J22")]
    joined = " ".join(markers)
    return "FIRST ARTICLE" in joined and "CHAR" in joined and "REQUIREMENT" in joined and "INSPECTION RESULT" in joined


def _copy_row_style(ws, source_row: int, target_row: int) -> None:
    for col in range(1, ws.max_column + 1):
        src = ws.cell(row=source_row, column=col)
        dst = ws.cell(row=target_row, column=col)
        if src.has_style:
            dst._style = copy.copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy.copy(src.alignment)
        if src.border:
            dst.border = copy.copy(src.border)
        if src.fill:
            dst.fill = copy.copy(src.fill)
        if src.font:
            dst.font = copy.copy(src.font)
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height or R3_BODY_ROW_HEIGHT


def _copy_merged_ranges_for_row(ws, source_row: int, target_row: int) -> None:
    existing = {str(rng) for rng in ws.merged_cells.ranges}
    for merged in list(ws.merged_cells.ranges):
        if merged.min_row == source_row and merged.max_row == source_row:
            row_delta = target_row - source_row
            ref = f"{get_column_letter(merged.min_col)}{merged.min_row + row_delta}:{get_column_letter(merged.max_col)}{merged.max_row + row_delta}"
            if ref not in existing:
                ws.merge_cells(ref)
                existing.add(ref)


def _ensure_list_sheets(wb) -> None:
    for sheet_name, values in LIST_SHEET_SPECS.items():
        ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.create_sheet(sheet_name)
        for row, value in enumerate(values, start=1):
            ws.cell(row=row, column=1).value = value


def _add_list_validation(ws, range_ref: str, formula1: str) -> None:
    validation = DataValidation(type="list", formula1=formula1, allow_blank=True)
    validation.error = "Choose a listed value or leave blank."
    validation.errorTitle = "Invalid selection"
    validation.prompt = "Choose from the list."
    validation.promptTitle = "EZ FAI"
    ws.add_data_validation(validation)
    validation.add(range_ref)


def _r3_inclusive_formula(row: int) -> str:
    return f'=IF(J{row}="","",IF(AND(J{row}>=C{row},J{row}<=E{row}),"X",""))'


def _required_r3_end_row(characteristic_count: int) -> int:
    return max(R3_TEMPLATE_END_ROW, R3_START_ROW + max(characteristic_count, 1) - 1)


def _ensure_r3_row_capacity(ws, characteristic_count: int) -> int:
    required_end_row = _required_r3_end_row(characteristic_count)
    if required_end_row <= R3_TEMPLATE_END_ROW:
        return R3_TEMPLATE_END_ROW
    extra_rows = required_end_row - R3_TEMPLATE_END_ROW
    insert_at = R3_TEMPLATE_END_ROW + 1
    ws.insert_rows(insert_at, amount=extra_rows)
    for row in range(insert_at, required_end_row + 1):
        _copy_row_style(ws, R3_TEMPLATE_END_ROW, row)
        _copy_merged_ranges_for_row(ws, R3_TEMPLATE_END_ROW, row)
    return required_end_row


def _reset_r3_rows(ws, start_row: int, end_row: int) -> None:
    for row in range(start_row, end_row + 1):
        ws.cell(row=row, column=R3_COLUMNS["Char Number"]).value = row - start_row + 1
        for header in ["Reference Location", "Requirement LSL", "Requirement Nominal", "Requirement USL", "Type", "Supplier Actual", "Supplier Yes", "Supplier No", "EZ Fabricating Actual", "Tooling Used", "Comments"]:
            ws.cell(row=row, column=R3_COLUMNS[header]).value = None
        ws.cell(row=row, column=R3_COLUMNS["In Spec"]).value = _r3_inclusive_formula(row)


def _apply_r3_editable_layout(ws, end_row: int) -> None:
    if ws.sheet_properties.pageSetUpPr is None:
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    else:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:N{end_row}"
    ws.print_title_rows = "1:23"
    ws.freeze_panes = "A24"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.scale = None
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.25
    ws.page_margins.bottom = 0.25
    ws.page_margins.header = 0.1
    ws.page_margins.footer = 0.1
    ws.sheet_view.view = "normal"
    ws.sheet_view.showGridLines = False
    for row in range(22, 24):
        ws.row_dimensions[row].height = R3_HEADER_ROW_HEIGHT
    for row in range(R3_START_ROW, end_row + 1):
        ws.row_dimensions[row].height = R3_BODY_ROW_HEIGHT
        for col in range(1, 15):
            cell = ws.cell(row=row, column=col)
            existing = cell.alignment or Alignment()
            cell.alignment = Alignment(horizontal=existing.horizontal or "center", vertical="center", text_rotation=existing.textRotation, wrap_text=True, shrink_to_fit=False, indent=existing.indent)


def _fill_r3_form(ws, characteristics: list[Any]) -> int:
    end_row = _ensure_r3_row_capacity(ws, len(characteristics))
    _reset_r3_rows(ws, R3_START_ROW, end_row)
    for offset, characteristic in enumerate(characteristics):
        row = R3_START_ROW + offset
        values = _row_values(characteristic)
        ws.cell(row=row, column=R3_COLUMNS["Char Number"]).value = offset + 1
        ws.cell(row=row, column=R3_COLUMNS["Reference Location"]).value = values["Reference Location"]
        ws.cell(row=row, column=R3_COLUMNS["Requirement LSL"]).value = values["Requirement LSL"]
        ws.cell(row=row, column=R3_COLUMNS["Requirement Nominal"]).value = values["Requirement Nominal"]
        ws.cell(row=row, column=R3_COLUMNS["Requirement USL"]).value = values["Requirement USL"]
        ws.cell(row=row, column=R3_COLUMNS["Type"]).value = values["Type"]
        ws.cell(row=row, column=R3_COLUMNS["EZ Fabricating Actual"]).value = None
        ws.cell(row=row, column=R3_COLUMNS["In Spec"]).value = _r3_inclusive_formula(row)
        ws.cell(row=row, column=R3_COLUMNS["Tooling Used"]).value = values["Tooling Used"]
        ws.cell(row=row, column=R3_COLUMNS["Comments"]).value = values["Comments"]
    _add_list_validation(ws, f"F{R3_START_ROW}:F{end_row}", "'CHARACTERISTICS'!$A$1:$A$33")
    _add_list_validation(ws, f"M{R3_START_ROW}:M{end_row}", "'TOOLING'!$A$1:$A$10")
    _add_list_validation(ws, f"H{R3_START_ROW}:I{end_row}", "'ATTRIBUTE'!$A$1:$A$2")
    _apply_r3_editable_layout(ws, end_row)
    return len(characteristics)


def _header_targets() -> dict[str, str]:
    targets: dict[str, str] = {}
    for canonical, aliases in HEADER_ALIASES.items():
        targets[_norm(canonical)] = canonical
        for alias in aliases:
            targets[_norm(alias)] = canonical
    return targets


def _find_generic_header_row(ws) -> tuple[int, dict[str, int]]:
    targets = _header_targets()
    best_row = 1
    best_mapping: dict[str, int] = {}
    for row in range(1, min(ws.max_row, 80) + 1):
        mapping: dict[str, int] = {}
        for col in range(1, min(ws.max_column, 80) + 1):
            value = _norm(ws.cell(row=row, column=col).value)
            if value in targets:
                mapping.setdefault(targets[value], col)
        if len(mapping) > len(best_mapping):
            best_row, best_mapping = row, mapping
    if len(best_mapping) < 4:
        best_row = 1
        best_mapping = {header: col for col, header in enumerate(GENERIC_HEADERS, start=1)}
        for header, col in best_mapping.items():
            ws.cell(row=best_row, column=col).value = header
    else:
        for header in GENERIC_HEADERS:
            if header not in best_mapping:
                col = ws.max_column + 1
                ws.cell(row=best_row, column=col).value = header
                best_mapping[header] = col
    return best_row, best_mapping


def _fill_generic(ws, characteristics: list[Any]) -> int:
    header_row, columns = _find_generic_header_row(ws)
    start_row = header_row + 1
    for offset, characteristic in enumerate(characteristics):
        row = start_row + offset
        if offset > 0:
            _copy_row_style(ws, start_row, row)
        values = _row_values(characteristic)
        for header, value in values.items():
            if header == "EZ Fabricating Actual":
                value = None
            ws.cell(row=row, column=columns[header]).value = value
        actual_col = get_column_letter(columns["EZ Fabricating Actual"])
        lsl_col = get_column_letter(columns["Requirement LSL"])
        usl_col = get_column_letter(columns["Requirement USL"])
        in_spec_col = columns["In Spec"]
        ws.cell(row=row, column=in_spec_col).value = f'=IF({actual_col}{row}="","",IF(AND({actual_col}{row}>={lsl_col}{row},{actual_col}{row}<={usl_col}{row}),"X",""))'
    return len(characteristics)


def template_row_capacity(template_path: str | Path) -> int | None:
    wb = load_workbook(template_path, read_only=True, data_only=False)
    wb.close()
    return None


def fill_fai_template(template_path: str | Path, characteristics: Iterable[Any], output_path: str | Path | None = None) -> Path:
    template_path = Path(template_path)
    characteristics = list(characteristics)
    if output_path is None:
        drawing_name = "FAI"
        if characteristics:
            metadata = _characteristic_metadata(characteristics[0])
            drawing_name = metadata.get("drawing_name", drawing_name)
        suffix = ".xlsm" if template_path.suffix.lower() == ".xlsm" else ".xlsx"
        output_path = template_path.with_name(f"{drawing_name}_FAI{suffix}")
    output_path = Path(output_path)
    if output_path.resolve() == template_path.resolve():
        raise ValueError("FAI Excel output cannot overwrite the original template.")
    keep_vba = template_path.suffix.lower() == ".xlsm"
    wb = load_workbook(template_path, keep_vba=keep_vba)
    _ensure_list_sheets(wb)
    ws = _pick_sheet(wb)
    if _is_r3_form(ws):
        _fill_r3_form(ws, characteristics)
    else:
        _fill_generic(ws, characteristics)
    wb.save(output_path)
    wb.close()
    return output_path
